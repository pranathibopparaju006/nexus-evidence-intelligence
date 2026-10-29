import os
import json
import logging
from datetime import datetime
from typing import Dict, Any

import PyPDF2
from docx import Document
from PIL import Image
import pytesseract

from app.core.database import SessionLocal
from app.models.database import DataSource

logger = logging.getLogger(__name__)


async def process_file(source_id: int, file_path: str):
    """Process uploaded file and extract content."""
    db = SessionLocal()
    source = None
    try:
        source = db.query(DataSource).filter(DataSource.id == source_id).first()
        if not source:
            return
        
        source.ingestion_status = "processing"
        db.commit()
        
        file_ext = os.path.splitext(file_path)[1].lower()
        
        if file_ext == '.pdf':
            content = extract_from_pdf(file_path)
        elif file_ext == '.docx':
            content = extract_from_docx(file_path)
        elif file_ext in ['.txt', '.md']:
            content = extract_from_text(file_path)
        elif file_ext in ['.png', '.jpg', '.jpeg', '.gif']:
            content = extract_from_image(file_path)
        elif file_ext == '.json':
            content = extract_from_json(file_path)
        elif file_ext == '.csv':
            content = extract_from_csv(file_path)
        else:
            content = {"text": "", "error": f"Unsupported file type: {file_ext}"}
        
        source.extracted_metadata = {
            "content": content,
            "processed_at": datetime.utcnow().isoformat(),
            "file_size": os.path.getsize(file_path)
        }
        source.ingestion_status = "completed"
        
        db.commit()
        
    except Exception:
        logger.exception("Document ingestion failed for source_id=%s", source_id)
        if source is not None:
            source.ingestion_status = "error"
            source.extracted_metadata = {"error": "Document processing failed. Check server logs for details."}
            db.commit()
    finally:
        db.close()


def extract_from_pdf(file_path: str) -> Dict[str, Any]:
    text = ""
    with open(file_path, 'rb') as file:
        reader = PyPDF2.PdfReader(file)
        for page in reader.pages:
            text += (page.extract_text() or "") + "\n"
    
    return {
        "text": text,
        "pages": len(reader.pages),
        "type": "pdf"
    }


def extract_from_docx(file_path: str) -> Dict[str, Any]:
    doc = Document(file_path)
    text = "\n".join([para.text for para in doc.paragraphs])
    
    return {
        "text": text,
        "paragraphs": len(doc.paragraphs),
        "type": "docx"
    }


def extract_from_text(file_path: str) -> Dict[str, Any]:
    with open(file_path, 'r', encoding='utf-8') as file:
        text = file.read()
    
    return {
        "text": text,
        "type": "text"
    }


def extract_from_image(file_path: str) -> Dict[str, Any]:
    image = Image.open(file_path)
    text = pytesseract.image_to_string(image)
    
    return {
        "text": text,
        "image_size": image.size,
        "type": "image"
    }


def extract_from_json(file_path: str) -> Dict[str, Any]:
    with open(file_path, 'r') as file:
        data = json.load(file)
    
    text = json.dumps(data, indent=2)
    
    return {
        "text": text,
        "structured_data": data,
        "type": "json"
    }


def extract_from_csv(file_path: str) -> Dict[str, Any]:
    import csv
    
    rows = []
    with open(file_path, 'r', encoding='utf-8') as file:
        reader = csv.reader(file)
        for row in reader:
            rows.append(row)
    
    text = "\n".join([", ".join(row) for row in rows])
    
    return {
        "text": text,
        "rows": len(rows),
        "type": "csv"
    }
